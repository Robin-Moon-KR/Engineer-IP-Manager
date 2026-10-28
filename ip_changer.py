import json
import os
import subprocess
import tkinter as tk
from tkinter import filedialog, messagebox, ttk
import openpyxl

CONFIG_FILE = "network_profiles.json"


class IPChangerApp:

  def __init__(self, root):
    self.root = root
    self.root.title("Equipment Setup IP Manager_ESIM")
    self.root.geometry("340x490")
    self.root.resizable(False, False)

    # 콤보박스 기본 좌측 정렬 설정
    self.style = ttk.Style()
    self.style.theme_use("default")
    self.style.configure(
        "Left.TCombobox", selectbackground="#0078d7", fieldbackground="white"
    )
    self.root.option_add("*TCombobox*Listbox.justify", "left")

    self.profiles = {}
    self.load_profiles()
    self.create_widgets()
    self.refresh_adapters()
    self.update_profile_listbox(list(self.profiles.keys()))

  def create_widgets(self):
    pad_options = {"padx": 10, "pady": 3}

    # 1. 어댑터 선택 프레임
    adapter_frame = tk.LabelFrame(
        self.root, text="1. Network Adaptor", font=("맑은 고딕", 9, "bold")
    )
    adapter_frame.pack(fill="x", **pad_options)

    self.adapter_combo = ttk.Combobox(
        adapter_frame,
        state="readonly",
        width=41,
        font=("맑은 고딕", 9),
        style="Left.TCombobox",
    )
    self.adapter_combo.pack(padx=10, pady=4)

    # 2. 저장된 프로필 검색 및 선택 프레임
    load_frame = tk.LabelFrame(
        self.root, text="2. Find & Select Profile", font=("맑은 고딕", 9, "bold")
    )
    load_frame.pack(fill="x", **pad_options)

    search_sub = tk.Frame(load_frame)
    search_sub.pack(fill="x", padx=10, pady=3)

    tk.Label(search_sub, text="Search:", font=("맑은 고딕", 9)).pack(
        side="left", padx=(0, 5)
    )

    self.search_var = tk.StringVar()
    self.search_var.trace_add("write", self.on_search_changed)

    self.search_entry = tk.Entry(
        search_sub, textvariable=self.search_var, width=19, font=("맑은 고딕", 9)
    )
    self.search_entry.pack(side="left", padx=(0, 5))

    tk.Button(
        search_sub,
        text="search Reset",
        command=self.reset_search,
        font=("맑은 고딕", 9),
    ).pack(side="left")

    # 실시간 검색 결과 리스트박스 (좌측 정렬)
    list_sub = tk.Frame(load_frame)
    list_sub.pack(fill="x", padx=10, pady=(0, 5))

    self.profile_listbox = tk.Listbox(
        list_sub,
        height=4,
        font=("맑은 고딕", 9),
        exportselection=False,
        justify="left",
    )
    self.profile_listbox.pack(side="left", fill="x", expand=True)
    self.profile_listbox.bind(
        "<<ListboxSelect>>", self.on_profile_selected_from_list
    )

    scrollbar = tk.Scrollbar(
        list_sub, orient="vertical", command=self.profile_listbox.yview
    )
    scrollbar.pack(side="right", fill="y")
    self.profile_listbox.config(yscrollcommand=scrollbar.set)

    # 3. 입력 정보 프레임
    input_frame = tk.LabelFrame(
        self.root,
        text="3. Equipment and Network Information",
        font=("맑은 고딕", 9, "bold"),
    )
    input_frame.pack(fill="x", **pad_options)

    self.entries = {}
    normal_fields = [
        ("Customer Name:", "cust_entry"),
        ("Location in building:", "loc_entry"),
        ("Device Name:", "dev_entry"),
    ]

    for label_text, attr_name in normal_fields:
      row_frame = tk.Frame(input_frame)
      row_frame.pack(fill="x", padx=10, pady=2)

      tk.Label(
          row_frame, text=label_text, width=18, anchor="w", font=("맑은 고딕", 9)
      ).pack(side="left")
      entry = tk.Entry(row_frame, width=24, font=("맑은 고딕", 9))
      entry.pack(side="left")
      self.entries[attr_name] = entry

    ip_fields = [
        ("IP Address:", "ip"),
        ("Subnet Mask:", "subnet"),
        ("Gateway:", "gateway"),
    ]

    self.ip_octets = {}
    for label_text, key_prefix in ip_fields:
      row_frame = tk.Frame(input_frame)
      row_frame.pack(fill="x", padx=10, pady=2)

      tk.Label(
          row_frame, text=label_text, width=18, anchor="w", font=("맑은 고딕", 9)
      ).pack(side="left")

      octet_container = tk.Frame(row_frame)
      octet_container.pack(side="left")

      octets = []
      for i in range(4):
        e = tk.Entry(
            octet_container, width=4, justify="center", font=("맑은 고딕", 9)
        )
        e.pack(side="left")
        octets.append(e)

        e.bind(
            "<FocusIn>",
            lambda event, entry=e: self.root.after(
                50, lambda: entry.select_range(0, tk.END)
            ),
        )
        e.bind(
            "<KeyPress>",
            lambda event, cur=i, o_list=octets: self.check_octet_arrow(
                event, cur, o_list
            ),
        )
        e.bind(
            "<KeyRelease>",
            lambda event, cur=i, o_list=octets: self.check_octet_jump(
                event, cur, o_list
            ),
        )

        if i < 3:
          tk.Label(
              octet_container, text=".", font=("맑은 고딕", 9, "bold")
          ).pack(side="left", padx=1)

      self.ip_octets[key_prefix] = octets

    # 4. 액션 버튼 프레임 (버튼들을 중앙 정렬로 배치)
    btn_frame = tk.Frame(self.root)
    btn_frame.pack(fill="x", padx=10, pady=6)

    row1_btn = tk.Frame(btn_frame)
    row1_btn.pack(anchor="center", pady=2)

    tk.Button(
        row1_btn,
        text="Save / Update",
        bg="#4CAF50",
        fg="white",
        font=("맑은 고딕", 9, "bold"),
        width=18,
        command=self.save_profile,
    ).pack(side="left", padx=2)
    tk.Button(
        row1_btn,
        text="Delete Info",
        bg="#f44336",
        fg="white",
        font=("맑은 고딕", 9, "bold"),
        width=17,
        command=self.delete_profile,
    ).pack(side="left", padx=2)

    row2_btn = tk.Frame(btn_frame)
    row2_btn.pack(anchor="center", pady=2)

    tk.Button(
        row2_btn,
        text="Export Excel",
        bg="#00838F",
        fg="white",
        font=("맑은 고딕", 9, "bold"),
        width=18,
        command=self.export_to_excel,
    ).pack(side="left", padx=2)
    tk.Button(
        row2_btn,
        text="Import Excel",
        bg="#00695C",
        fg="white",
        font=("맑은 고딕", 9, "bold"),
        width=17,
        command=self.import_from_excel,
    ).pack(side="left", padx=2)

    row3_btn = tk.Frame(btn_frame)
    row3_btn.pack(anchor="center", pady=2)

    tk.Button(
        row3_btn,
        text="Apply IP Set",
        bg="#2196F3",
        fg="white",
        font=("맑은 고딕", 9, "bold"),
        width=18,
        command=self.apply_network_setting,
    ).pack(side="left", padx=2)
    tk.Button(
        row3_btn,
        text="Apply DHCP Set",
        bg="#FF9800",
        fg="white",
        font=("맑은 고딕", 9, "bold"),
        width=17,
        command=self.apply_dhcp_setting,
    ).pack(side="left", padx=2)

    # 5. 하단 저작권 표기 (하단 여유 공간 중앙 정렬)
    footer_label = tk.Label(
        self.root,
        text="Copyright 2026. MOON_S. All rights reserved",
        font=("맑은 고딕", 8),
        fg="gray",
    )
    footer_label.pack(side="bottom", pady=5)

  def on_search_changed(self, *args):
    keyword = self.search_var.get().strip().lower()
    keys = list(self.profiles.keys())
    if keyword:
      filtered = [k for k in keys if keyword in k.lower()]
    else:
      filtered = keys
    self.update_profile_listbox(filtered)

  def update_profile_listbox(self, items):
    self.profile_listbox.delete(0, tk.END)
    for item in items:
      self.profile_listbox.insert(tk.END, item)

  def reset_search(self):
    self.search_var.set("")
    self.update_profile_listbox(list(self.profiles.keys()))
    self.search_entry.focus_set()

  def on_profile_selected_from_list(self, event):
    selection = self.profile_listbox.curselection()
    if selection:
      selected_key = self.profile_listbox.get(selection[0])
      if selected_key in self.profiles:
        data = self.profiles[selected_key]
        self.entries["cust_entry"].delete(0, tk.END)
        self.entries["cust_entry"].insert(0, data.get("customer", ""))

        self.entries["loc_entry"].delete(0, tk.END)
        self.entries["loc_entry"].insert(0, data.get("location", ""))

        self.entries["dev_entry"].delete(0, tk.END)
        self.entries["dev_entry"].insert(0, data.get("device", ""))

        self.set_ip_string("ip", data.get("ip", ""))
        self.set_ip_string("subnet", data.get("subnet", ""))
        self.set_ip_string("gateway", data.get("gateway", ""))

  def check_octet_arrow(self, event, idx, octet_list):
    if event.keysym == "Left":
      if idx > 0:
        octet_list[idx - 1].focus_set()
        return "break"
    elif event.keysym == "Right":
      if idx < len(octet_list) - 1:
        octet_list[idx + 1].focus_set()
        return "break"

  def check_octet_jump(self, event, idx, octet_list):
    val = octet_list[idx].get()
    if event.keysym in [
        "BackSpace",
        "Left",
        "Right",
        "Tab",
        "Shift_L",
        "Shift_R",
    ]:
      return

    if (
        ("." in val or len(val.replace(".", "")) >= 3)
        and idx < len(octet_list) - 1
    ):
      cleaned = val.replace(".", "").strip()
      octet_list[idx].delete(0, tk.END)
      octet_list[idx].insert(0, cleaned[:3])
      octet_list[idx + 1].focus_set()

  def get_ip_string(self, key_prefix):
    parts = [e.get().strip() for e in self.ip_octets[key_prefix]]
    if any(parts):
      return ".".join(parts)
    return ""

  def set_ip_string(self, key_prefix, ip_str):
    for e in self.ip_octets[key_prefix]:
      e.delete(0, tk.END)
    if ip_str:
      parts = ip_str.split(".")
      for i, part in enumerate(parts):
        if i < len(self.ip_octets[key_prefix]):
          self.ip_octets[key_prefix][i].insert(0, part)

  def clear_all_ip_octets(self):
    for key in self.ip_octets:
      for e in self.ip_octets[key]:
        e.delete(0, tk.END)

  def load_profiles(self):
    if os.path.exists(CONFIG_FILE):
      try:
        with open(CONFIG_FILE, "r", encoding="utf-8") as f:
          self.profiles = json.load(f)
      except Exception:
        self.profiles = {}

  def save_profiles_to_file(self):
    try:
      with open(CONFIG_FILE, "w", encoding="utf-8") as f:
        json.dump(self.profiles, f, ensure_ascii=False, indent=4)
    except Exception as e:
      messagebox.showerror("오류", f"설정 저장 실패: {e}")

  def refresh_adapters(self):
    try:
      cmd = 'powershell -Command "Get-NetAdapter | Where-Object {$_.Status -eq \\"Up\\"} | Select-Object -ExpandProperty Name"'
      result = subprocess.run(
          cmd,
          capture_output=True,
          text=True,
          shell=True,
          encoding="cp949",
          errors="ignore",
      )
      adapters = [line.strip() for line in result.stdout.splitlines() if line.strip()]
      if not adapters:
        adapters = ["이더넷"]
      self.adapter_combo["values"] = adapters
      self.adapter_combo.current(0)
    except Exception:
      self.adapter_combo["values"] = ["이더넷"]
      self.adapter_combo.current(0)

  def get_profile_key(self, cust, loc, dev):
    return f"[{cust}] {loc} - {dev}"

  def save_profile(self):
    cust = self.entries["cust_entry"].get().strip()
    loc = self.entries["loc_entry"].get().strip()
    dev = self.entries["dev_entry"].get().strip()

    ip = self.get_ip_string("ip")
    subnet = self.get_ip_string("subnet")
    gateway = self.get_ip_string("gateway")

    if not cust or not dev or not ip:
      messagebox.showwarning(
          "경고", "Customer Name, Device Name, IP Address는 필수 입력 항목입니다."
      )
      return

    key = self.get_profile_key(cust, loc, dev)
    self.profiles[key] = {
        "customer": cust,
        "location": loc,
        "device": dev,
        "ip": ip,
        "subnet": subnet,
        "gateway": gateway,
    }

    self.save_profiles_to_file()
    self.update_profile_listbox(list(self.profiles.keys()))
    messagebox.showinfo("성공", "프로필이 성공적으로 저장되었습니다.")

  def export_to_excel(self):
    """현재 저장된 프로필을 엑셀 파일로 내보내기"""
    if not self.profiles:
      messagebox.showwarning("경고", "내보낼 프로필 데이터가 없습니다.")
      return

    file_path = filedialog.asksaveasfilename(
        defaultextension=".xlsx",
        filetypes=[("Excel Files", "*.xlsx")],
        initialfile="network_profiles_backup.xlsx",
    )
    if not file_path:
      return

    try:
      wb = openpyxl.Workbook()
      ws = wb.active
      ws.title = "Network Profiles"

      # 헤더 작성
      headers = [
          "Customer Name",
          "Location in building",
          "Device Name",
          "IP Address",
          "Subnet Mask",
          "Gateway",
      ]
      ws.append(headers)

      # 데이터 작성
      for key, data in self.profiles.items():
        ws.append([
            data.get("customer", ""),
            data.get("location", ""),
            data.get("device", ""),
            data.get("ip", ""),
            data.get("subnet", ""),
            data.get("gateway", ""),
        ])

      wb.save(file_path)
      messagebox.showinfo("성공", f"엑셀 파일로 저장되었습니다:\n{file_path}")
    except Exception as e:
      messagebox.showerror("오류", f"엑셀 저장 중 오류 발생: {e}")

  def import_from_excel(self):
    """엑셀 파일을 불러와서 프로필 목록 병합 및 갱신"""
    file_path = filedialog.askopenfilename(
        filetypes=[("Excel Files", "*.xlsx")]
    )
    if not file_path:
      return

    try:
      wb = openpyxl.load_workbook(file_path)
      ws = wb.active

      added_count = 0
      # 2번째 줄(헤더 제외)부터 읽기
      for row in ws.iter_rows(min_row=2, values_only=True):
        cust, loc, dev, ip, subnet, gateway = (
            row[0],
            row[1],
            row[2],
            row[3],
            row[4],
            row[5],
        )

        if not cust or not dev or not ip:
          continue  # 필수 항목 누락 시 스킵

        cust = str(cust).strip()
        loc = str(loc).strip() if loc else ""
        dev = str(dev).strip()
        ip = str(ip).strip()
        subnet = str(subnet).strip() if subnet else ""
        gateway = str(gateway).strip() if gateway else ""

        key = self.get_profile_key(cust, loc, dev)
        self.profiles[key] = {
            "customer": cust,
            "location": loc,
            "device": dev,
            "ip": ip,
            "subnet": subnet,
            "gateway": gateway,
        }
        added_count += 1

      self.save_profiles_to_file()
      self.update_profile_listbox(list(self.profiles.keys()))
      messagebox.showinfo(
          "성공",
          f"엑셀에서 총 {added_count}개의 프로필을 성공적으로 불러왔습니다.",
      )
    except Exception as e:
      messagebox.showerror(
          "오류",
          f"엑셀 파일을 읽는 중 오류가 발생했습니다.\n올바른 양식인지 확인해"
          f" 주세요.\n상세: {e}",
      )

  def delete_profile(self):
    selection = self.profile_listbox.curselection()
    if not selection:
      messagebox.showwarning("경고", "삭제할 프로필을 목록에서 선택해주세요.")
      return

    selected_key = self.profile_listbox.get(selection[0])

    if messagebox.askyesno("확인", f"'{selected_key}' 프로필을 삭제하시겠습니까?"):
      if selected_key in self.profiles:
        del self.profiles[selected_key]
        self.save_profiles_to_file()
        self.update_profile_listbox(list(self.profiles.keys()))

        for entry in self.entries.values():
          entry.delete(0, tk.END)
        self.clear_all_ip_octets()

        messagebox.showinfo("완료", "프로필이 삭제되었습니다.")

  def apply_network_setting(self):
    adapter = self.adapter_combo.get()
    ip = self.get_ip_string("ip")
    subnet = self.get_ip_string("subnet")
    gateway = self.get_ip_string("gateway")

    if not adapter or not ip or not subnet:
      messagebox.showwarning(
          "경고",
          "어댑터 이름, IP 주소, 서브넷 마스크를 올바르게 입력(선택)해주세요.",
      )
      return

    if not messagebox.askyesno(
        "네트워크 변경 확인",
        f"어댑터 [{adapter}]의 설정을 고정 IP로 변경하시겠습니까?\n\nIP: {ip}\nSubnet:"
        f" {subnet}\nGateway: {gateway}",
    ):
      return

    if gateway:
      cmd = f'netsh interface ip set address name="{adapter}" static {ip} {subnet} {gateway}'
    else:
      cmd = (
          f'netsh interface ip set address name="{adapter}" static {ip}'
          f" {subnet}"
      )

    self.run_netsh_command(cmd)

  def apply_dhcp_setting(self):
    adapter = self.adapter_combo.get()
    if not adapter:
      messagebox.showwarning("경고", "대상 네트워크 어댑터를 선택해주세요.")
      return

    if not messagebox.askyesno(
        "DHCP 변경 확인",
        f"어댑터 [{adapter}]의 IP 설정을 **DHCP(자동 IP 받기)**로"
        " 변경하시겠습니까?",
    ):
      return

    cmd_ip = f'netsh interface ip set address name="{adapter}" source=dhcp'
    cmd_dns = f'netsh interface ip set dns name="{adapter}" source=dhcp'
    full_cmd = f"{cmd_ip} && {cmd_dns}"
    self.run_netsh_command(full_cmd)

  def run_netsh_command(self, cmd):
    try:
      result = subprocess.run(
          cmd,
          capture_output=True,
          text=True,
          shell=True,
          encoding="cp949",
          errors="ignore",
      )
      if result.returncode == 0:
        messagebox.showinfo(
            "성공", "네트워크 설정이 성공적으로 변경되었습니다."
        )
      else:
        error_msg = (
            result.stdout.strip()
            if result.stdout.strip()
            else result.stderr.strip()
        )
        messagebox.showerror(
            "오류 발생",
            f"설정 변경 실패:\n{error_msg}\n\n※ 반드시 프로그램을 '관리자 권한'으로"
            " 실행해 주세요.",
        )
    except Exception as e:
      messagebox.showerror("실행 오류", f"명령어 실행 중 오류가 발생했습니다: {e}")


if __name__ == "__main__":
  root = tk.Tk()
  app = IPChangerApp(root)
  root.mainloop()